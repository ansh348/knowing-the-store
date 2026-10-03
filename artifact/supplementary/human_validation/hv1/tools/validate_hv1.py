"""validate_hv1.py -- program R119, registration section 9.4: the return protocol and the completeness / integrity check.

    python validate_hv1.py hash  A2 <returned.xlsx>     sha256 of the returned file, WITHOUT opening it; appended to the ledger as AUX-HV1-RETURN-A2.
                                                        -> commit NUMBERS_HV1.md, then:
    python validate_hv1.py check A2 <returned.xlsx>     opens it: every locked cell unchanged against the key's hashes; label cells hold only allowed
                                                        values; counts filled / blank per block; writes results/labels_A2.json (labels and notes only).
    python validate_hv1.py hash  A1 <returned.xlsx>     -> commit, then
    python validate_hv1.py check A1 <returned.xlsx>     refuses to run unless AUX-HV1-RETURN-A2 is already in the COMMITTED ledger (git HEAD), or
                                                        --a2-absent is given, which records the pre-declared deviation (registration section 10).

    --dry-run <dir>   writes everything under <dir> and touches neither the ledger nor results/; for testing the tool on a synthetic fill.

The returned workbooks are never committed (registration section 0: A1's identity is on the ATTESTATION tab); they live under
workbooks/returned/, which is git-ignored. The labels file carries no name: the attestation is recorded as present / absent per field, and
the AI-tool line's answer text.
"""
from __future__ import annotations

import io
import json
import pathlib
import subprocess
import sys

import hv1_common as H
from hv1_common import J, ROOT, jdump, sha_text

RETURNED = H.WB_DIR / "returned"


def committed_ledger_has(key_name):
    try:
        out = subprocess.run(["git", "show", "HEAD:human_validation/hv1/NUMBERS_HV1.md"], cwd=ROOT, capture_output=True, text=True, encoding="utf-8")
        return out.returncode == 0 and key_name in out.stdout
    except Exception:
        return False


def cmd_hash(ann, path, dry):
    b = pathlib.Path(path).read_bytes(); sha = H.sha_bytes(b)
    line = (f"`{pathlib.Path(path).name}` ({len(b)} bytes) sha256 `{sha}`, hashed {H.now()} before the file was opened on this machine"
            + (" (A2's return hashed and to be committed before A1's is opened)." if ann == "A2" else "."))
    if dry:
        print(f"[dry-run] AUX-HV1-RETURN-{ann}: {line}")
    else:
        H.append_ledger(f"AUX-HV1-RETURN-{ann} (Phase 3)", line)
        print(f"AUX-HV1-RETURN-{ann} appended to the ledger: {sha}\n-> commit human_validation/hv1/NUMBERS_HV1.md now, then run `check`.")
    return sha


def cmd_check(ann, path, dry, a2_absent=False):
    from openpyxl import load_workbook
    key = J(H.KEY_DIR / "key_hv1.json"); assert key["meta"]["registration_sha256"] == H.REG_SHA
    wbk = key["workbooks"][ann]
    if ann == "A1" and not dry:
        if a2_absent:
            H.append_deviation("Registration sections 5.3 and 9.4: A2 did not return a workbook; A1's return opened without a committed AUX-HV1-RETURN-A2. "
                               "Kappa and sensitivity row (ii) print as dashes; the paper wording of the arrangement changes by ruling.")
        else:
            assert committed_ledger_has("AUX-HV1-RETURN-A2"), "A2's return is not hashed in the COMMITTED ledger; A1's file stays closed -- STOP"
        assert committed_ledger_has(f"AUX-HV1-RETURN-{ann}"), f"AUX-HV1-RETURN-{ann} is not in the committed ledger; run `hash` and commit first -- STOP"
    if ann == "A2" and not dry:
        assert committed_ledger_has("AUX-HV1-RETURN-A2"), "AUX-HV1-RETURN-A2 is not in the committed ledger; run `hash` and commit first -- STOP"
    sha = H.sha_file(path)
    wb = load_workbook(path)
    problems, checked = [], 0
    # 1. locked cells unchanged
    stored = wbk["locked_cells"]
    for ref, h in stored.items():
        sheet, coord = ref.split("!")
        if sheet not in wb.sheetnames:
            problems.append(f"sheet missing: {sheet}"); continue
        v = wb[sheet][coord].value; checked += 1
        if v is None or sha_text(str(v)) != h:
            # openpyxl round-trips ints as ints and floats as floats; the build hashed str(value) of what it wrote
            problems.append(f"locked cell changed: {ref} (now {str(v)[:60]!r})")
    expected_sheets = ["README", "PRACTICE", "ANSWERS"] + (["TRIPLES"] if ann == "A1" else []) + ["ATTESTATION"]
    if wb.sheetnames != expected_sheets:
        problems.append(f"sheet list {wb.sheetnames} != {expected_sheets}")
    # 2. labels
    labels = {"ANSWERS": {}, "TRIPLES": {}}
    counts = {}
    ws = wb["ANSWERS"]
    rowmap = {int(k): v for k, v in wbk["answers_rows"].items()}
    items = key["hva"]["items"]
    by_n = {}
    for r in range(2, ws.max_row + 1):
        n = ws.cell(r, 1).value
        if not isinstance(n, int):
            continue                                                   # the divider row
        lab = ws.cell(r, 6).value; note = ws.cell(r, 7).value
        lab = str(lab).strip() if lab not in (None, "") else None
        if lab is not None and lab not in H.LABELS_A:
            problems.append(f"ANSWERS row n={n}: label {lab!r} not in {H.LABELS_A}")
        uid = rowmap[n]; by_n[n] = uid
        labels["ANSWERS"][str(n)] = {"uid": uid, "label": lab, "note": (str(note) if note not in (None, "") else None)}
    for block in ("primary", "haiku"):
        rows = [v for v in labels["ANSWERS"].values() if items[v["uid"]]["block"] == block]
        counts[block] = {"rows": len(rows), "labelled": sum(v["label"] is not None for v in rows), "blank": sum(v["label"] is None for v in rows),
                         "unsure": sum(v["label"] == "UNSURE" for v in rows)}
    if set(by_n) != set(rowmap):
        problems.append(f"ANSWERS rows found {len(by_n)} != key {len(rowmap)}")
    if ann == "A1":
        wt = wb["TRIPLES"]; tmap = {int(k): v for k, v in wbk["triples_rows"].items()}
        for r in range(2, wt.max_row + 1):
            n = wt.cell(r, 2).value
            if not isinstance(n, int):
                continue
            lab = wt.cell(r, 6).value; note = wt.cell(r, 7).value
            lab = str(lab).strip() if lab not in (None, "") else None
            if lab is not None and lab not in H.LABELS_B:
                problems.append(f"TRIPLES row n={n}: label {lab!r} not in {H.LABELS_B}")
            labels["TRIPLES"][str(n)] = {**tmap[n], "a1_label": lab, "note": (str(note) if note not in (None, "") else None)}
        rows = list(labels["TRIPLES"].values())
        counts["triples"] = {"rows": len(rows), "labelled": sum(v["a1_label"] is not None for v in rows), "blank": sum(v["a1_label"] is None for v in rows)}
        if len(rows) != len(tmap):
            problems.append(f"TRIPLES rows found {len(rows)} != key {len(tmap)}")
    # 3. attestation: presence only (the name is never written to the repo); the AI-tool line's answer is recorded
    wa = wb["ATTESTATION"]
    att = {wa.cell(r, 1).value: wa.cell(r, 2).value for r in range(1, wa.max_row + 1) if wa.cell(r, 1).value}
    attestation = {"name_present": bool(att.get("Your name")), "date": str(att.get("Date finished") or "") or None,
                   "hours": str(att.get("Roughly how long it took (hours)") or "") or None,
                   "confusing_present": bool(att.get("Anything you found confusing")),
                   "no_ai_tool_answer": str(att.get("I did not use any AI tool while labelling (write YES)") or "") or None}
    out = {"meta": {"program": H.PROGRAM, "registration_sha256": H.REG_SHA, "annotator": ann, "returned_sha256": sha, "key_sha256": H.sha_file(H.KEY_DIR / "key_hv1.json"),
                    "validated": H.now(), "locked_cells_checked": checked, "problems": problems, "dry_run": bool(dry)},
           "counts": counts, "attestation": attestation, "labels": labels}
    outdir = pathlib.Path(dry) if dry else H.RES_DIR
    p = outdir / f"labels_{ann}.json"; jdump(out, p)
    verdict = "PASS" if not problems else "FAIL"
    summary = (f"`labels_{ann}.json` from the return sha256 `{sha}`: locked cells checked {checked}, **{verdict}** ({len(problems)} problems"
               + (": " + "; ".join(problems[:10]) if problems else "") + f"); counts {counts}; attestation name present {attestation['name_present']}, "
               f"hours {attestation['hours']}, AI-tool line {attestation['no_ai_tool_answer']!r}.")
    if dry:
        print(f"[dry-run] AUX-HV1-VALIDATE-{ann}: {summary}")
    else:
        H.append_ledger(f"AUX-HV1-VALIDATE-{ann} (Phase 3)", summary)
        print(f"validate {ann}: {verdict}; {counts} -> {p.relative_to(ROOT)}")
    if problems:
        print("\n".join("  " + q for q in problems[:20]))
        sys.exit(1)
    return out


if __name__ == "__main__":
    H.assert_registration()
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = None
    if "--dry-run" in sys.argv:
        dry = sys.argv[sys.argv.index("--dry-run") + 1]; args = [a for a in args if a != dry]
    cmd, ann, path = args[0], args[1], args[2]
    assert ann in ("A1", "A2")
    if cmd == "hash":
        cmd_hash(ann, path, dry)
    elif cmd == "check":
        cmd_check(ann, path, dry, a2_absent="--a2-absent" in sys.argv)
    else:
        sys.exit(f"unknown command {cmd}")
