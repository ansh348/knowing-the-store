"""build_hv1.py -- program R119: the frozen extract, the populations, the key, and the two blank workbooks.

    python build_hv1.py extract       -> inputs/hva_rows.json (the rows and per-row scores this program uses, with source digests)
    python build_hv1.py candidates    -> key/practice_candidates.json (the seeded practice candidates, for the builder's read)
    python build_hv1.py build         -> key/key_hv1.json, workbooks/blank/HV1_A1.xlsx, workbooks/blank/HV1_A2.xlsx, ledger sections

Everything follows REGISTRATION_HV1.md (asserted at start). No model call. HV-B versions are rebuilt with cf2_build_cell.py's
construction and asserted, version by version, against cf_cell.json's label and cf_indexes.json's v2 text.
"""
from __future__ import annotations

import collections
import copy
import datetime as dt
import random
import sys

import hv1_common as H
from hv1_common import J, ROOT, HV1, jdump, sha_text

# ----------------------------------------------------------------------------- the frozen extract
def extract():
    rows, scores, rec, byid = H.load_hva_sources()
    out = {"meta": {"program": H.PROGRAM, "registration_sha256": H.REG_SHA, "freeze_commit": H.FREEZE_COMMIT,
                    "sources": H.assert_inputs(), "when": H.now(),
                    "note": "GPT-4.1 mini (MINI) and Claude Haiku 4.5 (H45) answers on the 452 downstream views, before (Amendment 4 header) "
                            "and after (Amendment 5 retag) with the committed detector's per-row scores; stale / gold fields from run_ds.records."},
           "rows": {}}
    for model in ("MINI", "H45"):
        for key, rb in rows[(model, "before")].items():
            ra = rows[(model, "after")][key]; it = byid[rb["item_id"]]; rc = rec.get(rb["item_id"]) or {}
            out["rows"][f"{model}|{key}"] = {
                "model": model, "key": key, "item_id": rb["item_id"], "view_sha": rb["view_sha"], "persona": rb["persona"], "label": rb["label"],
                "state": rb["state"], "policies": rb["policies"], "question": rb["question"],
                "before": {"answer": rb["answer"], "source": rb["source"], "score": scores["before"][model][key]},
                "after": {"answer": ra["answer"], "source": ra["source"], "score": scores["after"][model][key]},
                "stale_field": rc.get("stale_field"), "gold_field": rc.get("gold_field"), "section": rc.get("section"), "intent": it.get("intent")}
    p = H.INPUTS_DIR / "hva_rows.json"; jdump(out, p)
    print(f"extract: {len(out['rows'])} rows -> {p.relative_to(ROOT)}  sha256 {H.sha_file(p)[:16]}")
    return out


# ----------------------------------------------------------------------------- HV-A populations
def hva_populations(X):
    R = X["rows"]
    mini = {r["key"]: r for r in R.values() if r["model"] == "MINI"}
    h45 = {r["key"]: r for r in R.values() if r["model"] == "H45"}
    scor = sorted(k for k, r in mini.items() if r["label"] == "PRESENT-NEGATIVE" and r["state"] in ("sup", "ret"))
    changed = [k for k in scor if r_src(mini[k]) == "changed"]
    unchanged = [k for k in scor if k not in set(changed)]
    assert len(scor) == 212 and len(changed) == 74 and len(unchanged) == 138, (len(scor), len(changed), len(unchanged))
    fb = sum(mini[k]["before"]["score"]["stale_use_tau_gold"] for k in changed); fa = sum(mini[k]["after"]["score"]["stale_use_tau_gold"] for k in changed)
    assert (fb, fa) == (37, 8), (fb, fa)
    assert not any(mini[k]["before"]["score"]["stale_use_tau_gold"] or mini[k]["after"]["score"]["stale_use_tau_gold"] for k in unchanged)
    identical = [k for k in changed if mini[k]["before"]["answer"] == mini[k]["after"]["answer"]]
    assert len(identical) == 20, len(identical)
    # controls: round robin over the six Table 4 policies, seed 11901 (registration 3.2)
    rng = random.Random(H.SEED_CONTROLS)
    pools = {p: sorted(k for k in unchanged if p in mini[k]["policies"]) for p in H.SIX_POLICIES}
    for p in H.SIX_POLICIES:
        rng.shuffle(pools[p])
    drawn, turns = [], []
    while len(drawn) < H.N_CONTROLS:
        progressed = False
        for p in H.SIX_POLICIES:
            while pools[p] and pools[p][-1] in drawn:
                pools[p].pop()
            if pools[p] and len(drawn) < H.N_CONTROLS:
                k = pools[p].pop(); drawn.append(k); turns.append(p); progressed = True
        assert progressed, "control pools exhausted"
    assert len(drawn) == 50 and len(set(drawn)) == 50
    # Haiku secondary block: flagged answers on the 74, before and after
    hb = [k for k in changed if h45[k]["before"]["score"]["stale_use_tau_gold"]]
    ha = [k for k in changed if h45[k]["after"]["score"]["stale_use_tau_gold"]]
    assert (len(hb), len(ha)) == (38, 12), (len(hb), len(ha))
    items = {}
    for k in changed:
        for cond in ("before", "after"):
            uid = f"MINI|{cond}|{k}"; r = mini[k]
            items[uid] = dict(block="primary", model="MINI", cond=cond, key=k, item_id=r["item_id"], persona=r["persona"], state=r["state"],
                              policies=r["policies"], flag_gold=bool(r[cond]["score"]["stale_use_tau_gold"]), flag_hi=bool(r[cond]["score"]["stale_use"]),
                              stale_cov=r[cond]["score"].get("stale_cov"), twin=f"MINI|{'after' if cond == 'before' else 'before'}|{k}",
                              identical_twin=k in identical, question=r["question"], stale=r["stale_field"], gold=r["gold_field"],
                              answer=r[cond]["answer"], abstains=H.is_abstention(r[cond]["answer"]))
    for k, p in zip(drawn, turns):
        r = mini[k]; uid = f"MINI|control|{k}"
        items[uid] = dict(block="primary", model="MINI", cond="control", key=k, item_id=r["item_id"], persona=r["persona"], state=r["state"],
                          policies=r["policies"], drawn_for_policy=p, flag_gold=False, flag_hi=bool(r["before"]["score"]["stale_use"]),
                          stale_cov=r["before"]["score"].get("stale_cov"), twin=None, identical_twin=False, question=r["question"],
                          stale=r["stale_field"], gold=r["gold_field"], answer=r["before"]["answer"], abstains=H.is_abstention(r["before"]["answer"]))
    for cond, ks in (("before", hb), ("after", ha)):
        for k in ks:
            r = h45[k]; uid = f"H45|{cond}|{k}"; other = f"H45|{'after' if cond == 'before' else 'before'}|{k}"
            items[uid] = dict(block="haiku", model="H45", cond=cond, key=k, item_id=r["item_id"], persona=r["persona"], state=r["state"],
                              policies=r["policies"], flag_gold=True, flag_hi=bool(r[cond]["score"]["stale_use"]), stale_cov=r[cond]["score"].get("stale_cov"),
                              twin=other if (k in hb and k in ha) else None, identical_twin=False, question=r["question"], stale=r["stale_field"],
                              gold=r["gold_field"], answer=r[cond]["answer"], abstains=H.is_abstention(r[cond]["answer"]))
    for uid, it in items.items():
        assert it["stale"] and it["gold"], uid
    cov = {p: sum(p in mini[k]["policies"] for k in drawn) for p in H.SIX_POLICIES}
    pop = {"changed_keys": changed, "identical_text_keys": identical, "unchanged_keys": unchanged, "control_keys": drawn, "control_turns": turns,
           "control_coverage": cov, "control_abstentions": sum(H.is_abstention(mini[k]["before"]["answer"]) for k in drawn),
           "flags": {"before": fb, "after": fa}, "haiku_before_flagged": hb, "haiku_after_flagged": ha,
           "counts": {"primary": sum(i["block"] == "primary" for i in items.values()), "haiku": sum(i["block"] == "haiku" for i in items.values()),
                      "by_state_74": dict(collections.Counter(mini[k]["state"] for k in changed)),
                      "personas_74": len({mini[k]["persona"] for k in changed}), "items_74": len({mini[k]["item_id"] for k in changed})}}
    return items, pop


def r_src(r):
    return "changed" if "amend5 call" in r["after"]["source"] else "unchanged"


# ----------------------------------------------------------------------------- HV-B: the construction, mirrored and asserted
class CF:
    """CF-2 (or CF-1, for the practice set) construction, one pair at a time, mirroring cf2_build_cell.build()."""

    def __init__(self, which):
        import cf2_build_cell as CB
        import build_cf_cell as B
        import e0_common as C
        from index_v2 import scrub
        self.CB, self.B, self.C, self.scrub = CB, B, C, scrub
        self.parse_dt, self.toks = CB.parse_dt, CB.toks
        d = ROOT / "r82_counterfactual"
        if which == "cf2":
            d2 = d / "cf2"
            self.pairs = J(d2 / "cf_pairs_cf2.json")["pairs"]
            self.qs = {r["pair_id"]: r for r in J(d2 / "cf_questions_cf2.json")["questions"]}
            self.pn = {r["pair_id"]: r for r in J(d2 / "cf_pn_texts_cf2.json")["records"]}
            self.names = J(d2 / "_names_cf2.json")
            self.cell = {it["item_id"]: it for it in J(d2 / "cf_cell.json")["items"]}
            self.idx = J(d2 / "cf_indexes.json")["indexes"]
        else:
            self.pairs = J(d / "cf_pairs.json")["pairs"]
            self.qs = {r["pair_id"]: r for r in J(d / "cf_questions.json")["questions"]}
            self.pn = {r["pair_id"]: r for r in J(d / "cf_pn_texts.json")["records"]}
            self.names = J(d / "_names.json")
            self.cell = {it["item_id"]: it for it in J(d / "cf_cell.json")["items"]}
            self.idx = J(d / "cf_indexes.json")["indexes"]
        self.ypaths = {p.name.replace(".yaml", "").split("_")[-1]: p for p in sorted(C.LME.glob("*/persona_*.yaml"))}
        self.banks = {}

    def bank(self, persona):
        if persona not in self.banks:
            self.banks[persona] = J(H.BANKS / f"{persona}.json")["entries"]
        return self.banks[persona]

    def slot_count(self, p):
        base = self.bank(p["persona"]); orig = next(e for e in base if e["id"] == p["entry_id"]); ss = self.scrub(orig["slot_key"][1])
        return sum(1 for e in base if e.get("slot_key") and self.scrub(e["slot_key"][1]) == ss)

    def in_window(self, p):
        return p["state"] == "EXPIRY" and self.parse_dt(p["event_time"]) < self.parse_dt(p["t_now"])

    def render(self, p):
        """The three versions of one pair: entries present, target-slot records, label; asserted against the committed cell and index."""
        B, C, CB = self.B, self.C, self.CB
        persona, state = p["persona"], p["state"]
        drop = self.toks(self.names[persona] or "")
        base = self.bank(persona); y0 = C.load_yaml(self.ypaths[persona])
        orig = next(e for e in base if e["id"] == p["entry_id"]); ss = self.scrub(orig["slot_key"][1])
        t_now, cutoff = self.parse_dt(p["t_now"]), self.parse_dt(p["cutoff"])
        newest = max(x for x in (self.parse_dt(e["created_at"]) for e in base if e["is_active"]) if x)
        JA = B.Join(persona, base, y0, t_now, cutoff, drop)
        goldA, firesA, expA = B.resolve_direct(JA, "stable_identity", {"fact": orig["text"], "session": orig["source_session"]})
        assert goldA, p["pair_id"]
        removedABS = {e["id"] for e in base if self.scrub(e.get("slot_key") and e["slot_key"][1]) == ss} | set(goldA)
        entB = copy.deepcopy(base); byidB = {e["id"]: e for e in entB}; yB = copy.deepcopy(y0)
        new_id = f"lb_r82_{p['pair_id'].split(':')[1]}"; sess_new = max(e.get("source_session") or 0 for e in base) + 1
        if state == "EXPIRY":
            past = (cutoff - dt.timedelta(days=30)).date().isoformat()
            byidB[orig["id"]]["event_time"] = past + "T00:00:00+00:00"
            yB.setdefault("expiring_logistics", []).append({"fact": orig["text"], "session": orig["source_session"], "event_date": past,
                                                            "status": "expired", "category": orig["category"]})
            secB, recB = "expiring_logistics", yB["expiring_logistics"][-1]
        else:
            new_text = self.pn[p["pair_id"]]["new_text"]
            conf = 0.9; thr = 0.95 if orig["category"] in CB.ACCUMULATE else 0.80
            entry = {"id": new_id, "text": new_text, "category": orig["category"], "category_weights": orig.get("category_weights"),
                     "slot_key": list(orig["slot_key"]), "event_time": None, "is_retracted": state == "RETRACTION", "confidence": conf,
                     "supersedes": orig["id"], "is_active": True, "created_at": newest.isoformat(), "source_session": sess_new,
                     "emotional_loading": 0.0, "numeric_value": False}
            fires_branch = conf >= thr
            if not fires_branch:
                entry["supersedes"] = None
            entB.append(entry); byidB[new_id] = entry
            if fires_branch:
                byidB[orig["id"]]["is_active"] = False
            if state == "SUPERSESSION":
                yB.setdefault("supersession_events", []).append({"slot": orig["slot_key"][1], "old_fact": orig["text"], "old_session": orig["source_session"],
                                                                 "new_fact": new_text, "new_session": sess_new, "category": orig["category"]})
                secB, recB = "supersession_events", yB["supersession_events"][-1]
            else:
                yB.setdefault("retractions", []).append({"original_fact": orig["text"], "original_session": orig["source_session"], "retraction": new_text,
                                                         "retraction_session": sess_new, "category": orig["category"]})
                secB, recB = "retractions", yB["retractions"][-1]
        JB = B.Join(persona, entB, yB, t_now, cutoff, drop)
        goldB, firesB, expB = B.resolve_direct(JB, secB, recB)
        out = {"pair_id": p["pair_id"], "persona": persona, "state": state, "intent": p["intent"], "question": self.qs[p["pair_id"]]["question"],
               "t_now": t_now.date().isoformat(), "cutoff": cutoff.date().isoformat(), "slot": ss, "category": orig["category"],
               "orig_id": orig["id"], "removedABS": sorted(removedABS), "versions": {}}
        for v in CB.VERSIONS:
            if v == "PN":
                Jx, removed, g, f_, x = JB, set(), goldB, firesB, expB
            else:
                Jx, g, f_, x = JA, goldA, firesA, expA
                removed = set() if v == "PRESENT" else removedABS
            lab, ev = B.label_of(g, f_, x, removed, p["intent"])
            assert lab == CB.WANT[v], (p["pair_id"], v, lab)
            key = f"{p['pair_id']}|{v}"
            assert self.cell[key]["label"] == lab and sorted(self.cell[key]["pn_members"]) == sorted(ev), key
            i2, i3, _ = B.build_index(Jx, removed)
            assert i2 == self.idx[key]["v2"], f"v2 index text differs from the committed build: {key}"
            present = [e for e in Jx.entries if e["id"] not in removed]
            target = [e for e in present if e.get("slot_key") and self.scrub(e["slot_key"][1]) == ss]
            out["versions"][v] = {"label": lab, "pn_members": sorted(ev), "n_entries": len(present),
                                  "target_records": [rec_view(e) for e in target]}
        # distractor pool: same category, other slot, a slot key, present in all three versions (registration 6.4 a-c; d is applied by the caller)
        out["distractor_pool"] = [rec_view(e) for e in base if e["category"] == orig["category"] and e["id"] not in removedABS
                                  and e.get("slot_key") and self.scrub(e["slot_key"][1]) != ss]
        return out


def rec_view(e):
    return {"id": e["id"], "text": e["text"], "slot": e["slot_key"][1] if e.get("slot_key") else None,
            "created": (e.get("created_at") or "")[:10] or None, "event": (e.get("event_time") or "")[:10] or None}


class Neighbour:
    """Registration 6.4 (d): cos(slot, target) under all-MiniLM-L6-v2 on the names with underscores as spaces, tau = 0.503637."""

    def __init__(self):
        from run_baselines import encoder
        from index_v2 import scrub
        self.enc, self.scrub, self.cache = encoder(), scrub, {}

    def vec(self, name):
        if name not in self.cache:
            self.cache[name] = self.enc.encode([name.replace("_", " ")], normalize_embeddings=True)[0]
        return self.cache[name]

    def cos(self, a, b):
        return float(self.vec(a) @ self.vec(b))

    def is_neighbour(self, slot, target):
        return self.cos(self.scrub(slot), self.scrub(target)) >= H.TAU_NEIGHBOUR


def hvb_population(cf, nb):
    rows = []
    for p in cf.pairs:
        n = cf.slot_count(p); w = cf.in_window(p)
        rows.append((p, n, w))
    excluded = {"slot_gt_10": collections.Counter(p["state"] for p, n, w in rows if n > H.MAX_SLOT_RECORDS),
                "expiry_in_window": collections.Counter(p["state"] for p, n, w in rows if n <= H.MAX_SLOT_RECORDS and w)}
    elig = [p for p, n, w in rows if n <= H.MAX_SLOT_RECORDS and not w]
    by_state = {s: sorted(p["pair_id"] for p in elig if p["state"] == s) for s in ("SUPERSESSION", "RETRACTION", "EXPIRY")}
    assert {s: len(v) for s, v in by_state.items()} == {"SUPERSESSION": 79, "RETRACTION": 15, "EXPIRY": 18}, {s: len(v) for s, v in by_state.items()}
    rng = random.Random(H.SEED_TRIPLES)
    chosen = []
    for s in ("SUPERSESSION", "RETRACTION", "EXPIRY"):
        ids = list(by_state[s]); rng.shuffle(ids); chosen += ids[:H.N_PER_STATE]
    order = sorted(chosen); rng.shuffle(order)           # presentation order of the 45 triples, the same seed
    byid = {p["pair_id"]: p for p in cf.pairs}
    triples, pools_short, neighbours_dropped = [], [], []
    rng_d = random.Random(H.SEED_DISTRACTORS)
    for pid in order:
        p = byid[pid]; R = cf.render(p)
        pool = [d for d in R["distractor_pool"] if not nb.is_neighbour(d["slot"], R["slot"])]
        neighbours_dropped.append(len(R["distractor_pool"]) - len(pool))
        pool.sort(key=lambda d: d["id"]); rng_d.shuffle(pool)
        distractors = pool[:H.N_DISTRACTORS]
        if len(distractors) < H.N_DISTRACTORS:
            pools_short.append((pid, len(distractors)))
        rng_t = random.Random(f"{H.SEED_WITHIN}:{pid}")
        vorder = list(cf.CB.VERSIONS); rng_t.shuffle(vorder)
        versions = []
        for v in vorder:
            recs = R["versions"][v]["target_records"] + distractors
            rng_t.shuffle(recs)
            versions.append({"version": v, "label": R["versions"][v]["label"], "pn_members": R["versions"][v]["pn_members"],
                             "records": recs, "n_target": len(R["versions"][v]["target_records"])})
        triples.append({"pair_id": pid, "persona": R["persona"], "state": R["state"], "intent": R["intent"], "question": R["question"],
                        "today": R["t_now"], "cutoff": R["cutoff"], "slot": R["slot"], "category": R["category"],
                        "distractor_ids": [d["id"] for d in distractors], "pool_size": len(pool), "versions": versions})
    pop = {"eligible_by_state": {s: len(v) for s, v in by_state.items()}, "excluded": {k: dict(v) for k, v in excluded.items()},
           "chosen": chosen, "presentation_order": order, "pools_under_six": pools_short,
           "neighbours_dropped_total": sum(neighbours_dropped), "n_versions": 3 * len(order)}
    return triples, pop


# ----------------------------------------------------------------------------- practice set
def practice_candidates(X):
    """Seeded candidates for the builder's read (registration 8): expiry answers of GPT-4.1 mini, Haiku answers on unchanged scorable views
    not among the 50 controls, and CF-1 pairs free of CF-2 passing the slot rule."""
    R = X["rows"]; mini = {r["key"]: r for r in R.values() if r["model"] == "MINI"}; h45 = {r["key"]: r for r in R.values() if r["model"] == "H45"}
    items, pop = hva_populations(X)
    controls = set(pop["control_keys"]); changed = set(pop["changed_keys"])
    rng = random.Random(H.SEED_PRACTICE)
    exp = sorted(k for k, r in mini.items() if r["label"] == "PRESENT-NEGATIVE" and r["state"] == "exp")
    seen, exp_c = set(), []
    rng.shuffle(exp)
    for k in exp:
        t = (mini[k]["item_id"], mini[k]["before"]["answer"])
        if t in seen:
            continue
        seen.add(t); exp_c.append(k)
    hk = sorted(k for k, r in h45.items() if r["label"] == "PRESENT-NEGATIVE" and r["state"] in ("sup", "ret") and k not in changed and k not in controls)
    rng.shuffle(hk)
    cf1 = CF("cf1"); nb = Neighbour()
    cf2 = CF("cf2")
    k2e = {(p["persona"], p["entry_id"]) for p in cf2.pairs}; k2s = {(p["persona"], cf2.scrub(p["slot_key"][1])) for p in cf2.pairs}
    free = [p for p in cf1.pairs if (p["persona"], p["entry_id"]) not in k2e and (p["persona"], cf1.scrub(p["slot_key"][1])) not in k2s]
    assert len(free) == 23, len(free)
    free = [p for p in free if cf1.slot_count(p) <= H.MAX_SLOT_RECORDS and not cf1.in_window(p)]
    ids = sorted(p["pair_id"] for p in free); rng.shuffle(ids)
    byid = {p["pair_id"]: p for p in cf1.pairs}
    plan = ["PRESENT", "PN", "ABSENT", "PRESENT", "PN"]
    stores, rng_d = [], random.Random(H.SEED_DISTRACTORS + 1)
    for pid, v in zip(ids[:5], plan):
        R = cf1.render(byid[pid])
        pool = [d for d in R["distractor_pool"] if not nb.is_neighbour(d["slot"], R["slot"])]
        pool.sort(key=lambda d: d["id"]); rng_d.shuffle(pool); dis = pool[:H.N_DISTRACTORS]
        rng_t = random.Random(f"{H.SEED_WITHIN}:{pid}:{v}")
        recs = R["versions"][v]["target_records"] + dis; rng_t.shuffle(recs)
        stores.append({"pair_id": pid, "version": v, "label": R["versions"][v]["label"], "state": R["state"], "question": R["question"],
                       "today": R["t_now"], "records": recs, "n_target": len(R["versions"][v]["target_records"])})
    def ans(model, k, cond="before"):
        r = (mini if model == "MINI" else h45)[k]
        stale = r["gold_field"] if r["state"] == "exp" else r["stale_field"]; gold = "date passed" if r["state"] == "exp" else r["gold_field"]
        return {"model": model, "key": k, "state": r["state"], "question": r["question"], "stale": stale, "gold": gold, "answer": r[cond]["answer"],
                "flag_gold": bool(r[cond]["score"]["stale_use_tau_gold"])}
    cands = {"meta": {"seed": H.SEED_PRACTICE, "rule": "seven expiry answers of GPT-4.1 mini (distinct texts) and three Haiku 4.5 answers on unchanged "
                      "scorable views outside the controls, in seeded order; if a label is missing among the ten, the last drawn items are replaced by "
                      "the first later candidates carrying it (registration 8). The builder writes label and why into key/practice_labels.json.",
                      "cf1_free_pairs": 23, "cf1_free_passing_slot_rule": len(free), "when": H.now()},
             "answers_expiry": [ans("MINI", k) for k in exp_c], "answers_haiku": [ans("H45", k) for k in hk[:12]], "stores": stores}
    p = H.KEY_DIR / "practice_candidates.json"; jdump(cands, p)
    print(f"practice candidates: {len(exp_c)} expiry, {len(hk[:12])} haiku, {len(stores)} stores -> {p.relative_to(ROOT)}")
    return cands


def practice_set(cands, labels):
    """Apply registration 8's rule with the builder's labels: seven expiry + three Haiku, every label of the four present."""
    L = labels["answers"]                      # key -> {"label", "why"}
    exp = [a for a in cands["answers_expiry"]]; hk = [a for a in cands["answers_haiku"]]
    chosen = exp[:7] + hk[:3]
    later = exp[7:] + hk[3:]
    def labels_of(xs): return {L[a["model"] + "|" + a["key"]]["label"] for a in xs}
    missing = [l for l in ("ASSERTS", "MENTIONS", "CLEAN", "ABSTAINS") if l not in labels_of(chosen)]
    replaced = []
    for l in missing:
        cand = None
        for a in later:                                   # the first later candidate carrying the label; every scanned candidate must be labelled
            lab = L[a["model"] + "|" + a["key"]]["label"]
            if lab == l:
                cand = a; break
        assert cand is not None, f"no later practice candidate carries {l}"
        # the last drawn item whose label occurs more than once is replaced, so that no label is lost (the registration's "last drawn items")
        counts = collections.Counter(L[a["model"] + "|" + a["key"]]["label"] for a in chosen)
        idx = max(i for i, a in enumerate(chosen) if counts[L[a["model"] + "|" + a["key"]]["label"]] > 1)
        gone = chosen.pop(idx); chosen.append(cand); later.remove(cand); replaced.append((gone["model"] + "|" + gone["key"], cand["model"] + "|" + cand["key"], l))
    assert len(chosen) == 10 and len(labels_of(chosen) & {"ASSERTS", "MENTIONS", "CLEAN", "ABSTAINS"}) == 4
    answers = [{**a, "label": L[a["model"] + "|" + a["key"]]["label"], "why": L[a["model"] + "|" + a["key"]]["why"]} for a in chosen]
    stores = [{**s, "why": labels["stores"][s["pair_id"] + "|" + s["version"]]["why"]} for s in cands["stores"]]
    return {"answers": answers, "stores": stores, "replaced": replaced}


# ----------------------------------------------------------------------------- orders
def constrained_order(uids, twin_of, seed):
    """Seeded shuffle; the two answers of one view at least MIN_GAP rows apart; redrawn with seed + 1000 k until it holds (registration 2)."""
    k = 0
    while True:
        rng = random.Random(seed + 1000 * k); order = sorted(uids); rng.shuffle(order)
        pos = {u: i for i, u in enumerate(order)}
        ok = all(abs(pos[u] - pos[t]) >= H.MIN_GAP for u, t in twin_of.items() if t in pos)
        if ok:
            return order, k
        k += 1
        assert k < 200000, "spacing constraint unsatisfiable"


# ----------------------------------------------------------------------------- the README
GUIDE_COMMON = [
    ("h", "Thank you for doing this. Please read this page once before you start. It takes about five minutes."),
    ("p", "You will label {WHAT}. Do the PRACTICE tab first. It has the correct answers next to it so you can check yourself. Then do {ORDER}."),
    ("b", "Please do not use any AI tool to help you label. Your own judgment is the point of this study."),
    ("h", "Tab ANSWERS (about 2 hours, 198 rows)"),
    ("p", "Each row has a question a user asked, an old value that is no longer true, the value that is true now, and an answer written by an assistant. Pick one label:"),
    ("l", "ASSERTS: the answer treats the old value as true now. Example: the user cancelled a trip, and the answer tells them what to pack for it."),
    ("l", "MENTIONS: the old value appears, but the answer says it is in the past, changed or cancelled, or explains why it no longer holds. Example: \"You used to prefer tea, and now you prefer coffee.\""),
    ("l", "CLEAN: the old value does not appear at all, and the answer does answer the question."),
    ("l", "ABSTAINS: the answer refuses, says it does not know, or gives nothing usable. Example: \"I don't have enough information about that.\""),
    ("l", "UNSURE: you really cannot decide. Use this rarely and say why in the notes column."),
    ("p", "Judge only what the answer says. Do not judge whether the answer is polite, complete or well written."),
    ("p", "The 50 rows below the grey divider at the bottom of the sheet are optional: do them only if you have time."),
]
GUIDE_TRIPLES = [
    ("h", "Tab TRIPLES (about 1.5 hours, 45 groups of 3)"),
    ("p", "Each row has a question and the notes a memory holds about one topic, with dates. Pick one label:"),
    ("l", "PRESENT: the memory has what is needed to answer, and nothing in it would mislead."),
    ("l", "PN: the memory holds something on this topic that must not be used for this question, such as a cancelled plan, a replaced value, or an event whose date has passed. Pick PN even if the right information is also there."),
    ("l", "ABSENT: the memory has nothing about this topic."),
    ("p", "Read the question carefully. A question about how something changed is allowed to use the old value, so that case is PRESENT. A question about what is true now is not allowed to use it, so that case is PN."),
    ("p", "Some notes will be about neighbouring topics; judge only the ones that bear on the question."),
]
GUIDE_GENERAL = [
    ("h", "General"),
    ("p", "Work in the order given. Do not sort or filter the sheet. You can stop and continue any time. When you finish, fill in the ATTESTATION tab (your name, the date, roughly how long it took, and anything you found confusing) and send the file back."),
    ("p", "If something seems broken, write it in the notes column and keep going."),
]


def guide(annotator):
    what, order = ("two things", "ANSWERS, then TRIPLES") if annotator == "A1" else ("one thing", "ANSWERS")
    rows = [(k, t.replace("{WHAT}", what).replace("{ORDER}", order)) for k, t in GUIDE_COMMON]
    if annotator == "A1":
        rows += GUIDE_TRIPLES
    return rows + GUIDE_GENERAL


def guide_text(annotator):
    return "\n".join(("- " if k == "l" else "") + t for k, t in guide(annotator))


# ----------------------------------------------------------------------------- workbooks
def records_cell(recs):
    lines = []
    for r in recs:
        tail = f"recorded {r['created']}" if r["created"] else "recorded: no date"
        if r["event"]:
            tail += f"; event date {r['event']}"
        lines.append(f"• {r['text']}  ({tail})")
    return "\n".join(lines)


def build_workbook(annotator, order_primary, order_haiku, items, triples, practice, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); locked = {}
    wrap = Alignment(wrap_text=True, vertical="top"); bold = Font(bold=True); grey = PatternFill("solid", fgColor="DDDDDD")
    unlocked = Protection(locked=False)

    def put(ws, coord, value, lock=True, **fmt):
        c = ws[coord]; c.value = value; c.alignment = wrap
        for k, v in fmt.items():
            setattr(c, k, v)
        if lock:
            if value not in (None, ""):
                locked[f"{ws.title}!{coord}"] = sha_text(str(value))
        else:
            c.protection = unlocked
        return c

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # README
    ws = wb.active; ws.title = "README"; widths(ws, [110])
    r = 1
    for kind, text in guide(annotator):
        c = put(ws, f"A{r}", ("• " + text) if kind == "l" else text)
        if kind == "h":
            c.font = bold
        if kind == "b":
            c.font = bold
        r += 1
    ws.protection.sheet = True

    dv_a = DataValidation(type="list", formula1='"' + ",".join(H.LABELS_A) + '"', allow_blank=True, showErrorMessage=True)
    dv_b = DataValidation(type="list", formula1='"' + ",".join(H.LABELS_B) + '"', allow_blank=True, showErrorMessage=True)

    # PRACTICE
    ws = wb.create_sheet("PRACTICE"); widths(ws, [5, 40, 30, 30, 60, 14, 60])
    for j, h in enumerate(["n", "Question", "Old value (no longer true)", "True now", "Answer", "KEY label", "Why"], 1):
        put(ws, f"{get_column_letter(j)}1", h, font=bold)
    for i, a in enumerate(practice["answers"], 1):
        put(ws, f"A{i + 1}", i); put(ws, f"B{i + 1}", a["question"]); put(ws, f"C{i + 1}", a["stale"]); put(ws, f"D{i + 1}", a["gold"])
        put(ws, f"E{i + 1}", a["answer"]); put(ws, f"F{i + 1}", a["label"], font=bold); put(ws, f"G{i + 1}", a["why"])
    if annotator == "A1":
        r0 = len(practice["answers"]) + 3
        for j, h in enumerate(["n", "Today", "Question", "Memory notes", "KEY label", "Why"], 1):
            put(ws, f"{get_column_letter(j)}{r0}", h, font=bold)
        ws.column_dimensions["D"].width = 90
        for i, s in enumerate(practice["stores"], 1):
            put(ws, f"A{r0 + i}", i); put(ws, f"B{r0 + i}", s["today"]); put(ws, f"C{r0 + i}", s["question"])
            put(ws, f"D{r0 + i}", records_cell(s["records"])); put(ws, f"E{r0 + i}", s["label"], font=bold); put(ws, f"F{r0 + i}", s["why"])
    ws.freeze_panes = "A2"; ws.protection.sheet = True

    # ANSWERS
    ws = wb.create_sheet("ANSWERS"); widths(ws, [5, 42, 34, 34, 70, 14, 40])
    for j, h in enumerate(["n", "Question", "Old value (no longer true)", "True now", "Answer", "Label", "Notes"], 1):
        put(ws, f"{get_column_letter(j)}1", h, font=bold)
    ws.add_data_validation(dv_a)
    rowmap = {}; r = 2
    for n, uid in enumerate(order_primary, 1):
        it = items[uid]
        put(ws, f"A{r}", n); put(ws, f"B{r}", it["question"]); put(ws, f"C{r}", it["stale"]); put(ws, f"D{r}", it["gold"]); put(ws, f"E{r}", it["answer"])
        put(ws, f"F{r}", None, lock=False); put(ws, f"G{r}", None, lock=False); dv_a.add(f"F{r}")
        rowmap[n] = uid; r += 1
    put(ws, f"A{r}", "Optional block: the rows below are extra. Do them only if you have time.", font=bold)
    for j in range(1, 8):
        ws[f"{get_column_letter(j)}{r}"].fill = grey
    r += 1
    for n, uid in enumerate(order_haiku, len(order_primary) + 1):
        it = items[uid]
        put(ws, f"A{r}", n); put(ws, f"B{r}", it["question"]); put(ws, f"C{r}", it["stale"]); put(ws, f"D{r}", it["gold"]); put(ws, f"E{r}", it["answer"])
        put(ws, f"F{r}", None, lock=False); put(ws, f"G{r}", None, lock=False); dv_a.add(f"F{r}")
        rowmap[n] = uid; r += 1
    ws.freeze_panes = "A2"; ws.protection.sheet = True

    # TRIPLES (A1 only)
    triplemap = {}
    if annotator == "A1":
        ws = wb.create_sheet("TRIPLES"); widths(ws, [7, 5, 12, 42, 100, 12, 40])
        for j, h in enumerate(["Group", "n", "Today", "Question", "Memory notes", "Label", "Notes"], 1):
            put(ws, f"{get_column_letter(j)}1", h, font=bold)
        ws.add_data_validation(dv_b)
        r, n = 2, 1
        for g, t in enumerate(triples, 1):
            for v in t["versions"]:
                put(ws, f"A{r}", g); put(ws, f"B{r}", n); put(ws, f"C{r}", t["today"]); put(ws, f"D{r}", t["question"]); put(ws, f"E{r}", records_cell(v["records"]))
                put(ws, f"F{r}", None, lock=False); put(ws, f"G{r}", None, lock=False); dv_b.add(f"F{r}")
                triplemap[n] = {"pair_id": t["pair_id"], "version": v["version"], "label": v["label"], "group": g}
                r += 1; n += 1
        ws.freeze_panes = "A2"; ws.protection.sheet = True

    # ATTESTATION
    ws = wb.create_sheet("ATTESTATION"); widths(ws, [40, 80])
    for i, q in enumerate(["Your name", "Date finished", "Roughly how long it took (hours)", "Anything you found confusing",
                           "I did not use any AI tool while labelling (write YES)"], 1):
        put(ws, f"A{i}", q, font=bold); put(ws, f"B{i}", None, lock=False)
    ws.protection.sheet = True

    out_path.parent.mkdir(parents=True, exist_ok=True); wb.save(out_path)
    return rowmap, triplemap, locked


# ----------------------------------------------------------------------------- build
def build():
    H.assert_registration(); digests = H.assert_inputs()
    X = J(H.INPUTS_DIR / "hva_rows.json")
    assert X["meta"]["registration_sha256"] == H.REG_SHA
    items, popA = hva_populations(X)
    cf = CF("cf2"); nb = Neighbour()
    triples, popB = hvb_population(cf, nb)
    cands = J(H.KEY_DIR / "practice_candidates.json"); labels = J(H.KEY_DIR / "practice_labels.json")
    practice = practice_set(cands, labels)
    primary = sorted(u for u, it in items.items() if it["block"] == "primary"); haiku = sorted(u for u, it in items.items() if it["block"] == "haiku")
    twins = {u: it["twin"] for u, it in items.items() if it["twin"]}
    orders, ks = {}, {}
    for ann, s_p, s_h in (("A1", H.SEED_ORDER_A1, H.SEED_HAIKU_A1), ("A2", H.SEED_ORDER_A2, H.SEED_HAIKU_A2)):
        op, kp = constrained_order(primary, {u: t for u, t in twins.items() if u in set(primary)}, s_p)
        oh, kh = constrained_order(haiku, {u: t for u, t in twins.items() if u in set(haiku)}, s_h)
        orders[ann] = (op, oh); ks[ann] = {"primary": kp, "haiku": kh}
    key = {"meta": {"program": H.PROGRAM, "registration_sha256": H.REG_SHA, "freeze_commit": H.FREEZE_COMMIT, "inputs": digests,
                    "extract_sha256": H.sha_file(H.INPUTS_DIR / "hva_rows.json"), "when": H.now(),
                    "seeds": {"controls": H.SEED_CONTROLS, "order_A1": H.SEED_ORDER_A1, "order_A2": H.SEED_ORDER_A2, "triples": H.SEED_TRIPLES,
                              "distractors": H.SEED_DISTRACTORS, "within": H.SEED_WITHIN, "haiku_A1": H.SEED_HAIKU_A1, "haiku_A2": H.SEED_HAIKU_A2,
                              "practice": H.SEED_PRACTICE}, "spacing_redraws_k": ks},
           "hva": {"population": popA, "items": items}, "hvb": {"population": popB, "triples": triples},
           "practice": practice, "workbooks": {}}
    wbdir = H.WB_DIR / "blank"
    for ann in ("A1", "A2"):
        op, oh = orders[ann]
        path = wbdir / f"HV1_{ann}.xlsx"
        rowmap, triplemap, locked = build_workbook(ann, op, oh, items, triples, practice, path)
        key["workbooks"][ann] = {"path": path.relative_to(ROOT).as_posix(), "answers_rows": rowmap, "triples_rows": triplemap, "locked_cells": locked,
                                 "readme_sha256": sha_text(guide_text(ann))}
    for ann in ("A1", "A2"):
        key["workbooks"][ann]["sha256_blank"] = H.sha_file(ROOT / key["workbooks"][ann]["path"])
    kp = H.KEY_DIR / "key_hv1.json"; jdump(key, kp)
    ksha = H.sha_file(kp)
    # ledger
    H.append_ledger("AUX-HV1-REG, AUX-HV1-INPUTS (Phase 2, build)",
                    f"Registered-text sha256 `{H.REG_SHA}`, frozen at `{H.FREEZE_COMMIT}` (2026-09-19), asserted by `build_hv1.py` at start. "
                    f"Records commit of the three Amendment 4 files: `bf1dafb`. Inputs as pinned in registration section 12, digests found "
                    f"(working bytes; LF-normalised where they differ):\n\n" +
                    "\n".join(f"- `{k}`: `{v.get('sha256', v.get('sha256_combined'))[:16]}`" + (f" (LF `{v['sha256_lf'][:16]}`)" if v.get("sha256_lf") and v["sha256_lf"] != v["sha256"] else "")
                              for k, v in digests.items()) +
                    f"\n\nFrozen extract `human_validation/hv1/inputs/hva_rows.json` sha256 `{key['meta']['extract_sha256']}` ({len(X['rows'])} rows).")
    pa = popA
    H.append_ledger("AUX-HV1-POP-A, AUX-HV1-CONTROLS, AUX-HV1-POP-A2 (Phase 2, build)",
                    f"**POP-A**: the 74 changed scorable views (registration 3.1): by state {pa['counts']['by_state_74']}, {pa['counts']['items_74']} obligations, "
                    f"{pa['counts']['personas_74']} personas; 148 answers; detector flags at 0.50: {pa['flags']['before']} before, {pa['flags']['after']} after, "
                    f"45 flagged answers; identical-text twins {len(pa['identical_text_keys'])}. The keys are in `key/key_hv1.json` (`hva.population.changed_keys`).\n\n"
                    f"**CONTROLS**: 50 unflagged answers from the 138 unchanged scorable views, seed {H.SEED_CONTROLS}, round robin in Table 4 order; "
                    f"turns {dict(collections.Counter(pa['control_turns']))}; views drawn serving each policy {pa['control_coverage']}; "
                    f"abstentions among them {pa['control_abstentions']}. Primary rows {pa['counts']['primary']}.\n\n"
                    f"**POP-A2** (optional block): Haiku 4.5 flagged answers on the 74 views, {len(pa['haiku_before_flagged'])} before + {len(pa['haiku_after_flagged'])} after = "
                    f"{pa['counts']['haiku']} rows.\n\n"
                    f"Spacing redraws k (registration 2): {ks}.")
    pb = popB
    H.append_ledger("AUX-HV1-POP-B, AUX-HV1-DISTRACTORS (Phase 2, build)",
                    f"Governing text: the registration as frozen (no Amendment 1). Eligible pairs by state {pb['eligible_by_state']}; excluded: "
                    f"more than ten slot records {pb['excluded']['slot_gt_10']}, expiry event inside the 14-day window {pb['excluded']['expiry_in_window']}. "
                    f"Drawn 15 per state, seed {H.SEED_TRIPLES}: 45 triples, {pb['n_versions']} versions; every version re-derived with cf2_build_cell's construction "
                    f"and asserted equal to `cf_cell.json`'s label and pn_members and to `cf_indexes.json`'s v2 text. Presentation order and the within-triple "
                    f"orders are in `key/key_hv1.json` (`hvb.triples`).\n\n"
                    f"**DISTRACTORS**: six per pair from the same category under other slot keys, present in all three versions, outside the own-name "
                    f"neighbourhood (tau {H.TAU_NEIGHBOUR}); records dropped as neighbours over the 45 pools: {pb['neighbours_dropped_total']}; "
                    f"pools under six: {pb['pools_under_six'] or 'none'}.")
    H.append_ledger("AUX-HV1-PRACTICE (Phase 2, build)",
                    f"Ten answers ({sum(a['model'] == 'MINI' for a in practice['answers'])} GPT-4.1 mini expiry, {sum(a['model'] == 'H45' for a in practice['answers'])} Haiku 4.5) "
                    f"with key labels {dict(collections.Counter(a['label'] for a in practice['answers']))}; replacements under the missing-label rule: "
                    f"{practice['replaced'] or 'none'}. Five CF-1 store versions ({[s['version'] for s in practice['stores']]}) from the "
                    f"{cands['meta']['cf1_free_passing_slot_rule']} of the 23 CF-1 pairs free of CF-2 that pass the slot rule. Key labels are the builder's read; "
                    f"the practice set is discarded and never scored.")
    H.append_ledger("AUX-HV1-WORKBOOK-A1, AUX-HV1-WORKBOOK-A2, AUX-HV1-KEY (Phase 2, build; recorded before anything is sent)",
                    "\n".join(f"- `{key['workbooks'][a]['path']}` sha256 `{key['workbooks'][a]['sha256_blank']}` (README sha256 `{key['workbooks'][a]['readme_sha256'][:16]}`)"
                              for a in ("A1", "A2")) +
                    f"\n- `human_validation/hv1/key/key_hv1.json` sha256 `{ksha}` (never sent).")
    if any(v > 0 for a in ks.values() for v in a.values()):
        H.append_deviation(f"Registration section 2: the seeded shuffle was redrawn for the spacing constraint; k = {ks} (seed + 1000 k). Recorded, no other change.")
    if pb["pools_under_six"]:
        H.append_deviation(f"Registration section 6.4: distractor pools under six: {pb['pools_under_six']}.")
    print(f"build: key sha256 {ksha[:16]}; A1 {key['workbooks']['A1']['sha256_blank'][:16]}; A2 {key['workbooks']['A2']['sha256_blank'][:16]}; redraws {ks}")
    return key


if __name__ == "__main__":
    H.assert_registration()
    cmd = sys.argv[1] if len(sys.argv) > 1 else "build"
    if cmd == "extract":
        H.assert_inputs(); extract()
    elif cmd == "candidates":
        H.assert_inputs(); practice_candidates(J(H.INPUTS_DIR / "hva_rows.json"))
    elif cmd == "build":
        build()
    else:
        sys.exit(f"unknown command {cmd}")
