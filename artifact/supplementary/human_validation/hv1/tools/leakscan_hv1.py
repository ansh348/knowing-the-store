"""leakscan_hv1.py -- program R119, registration section 9.2: read every cell of both blank workbooks.

Hard fail on any forbidden token in a cell outside the content columns; content-column occurrences are listed for review, and the
condition-revealing ones among them (a policy name, retag, oracle, gated, a model name) are hard fails too. Writes results/leakscan_hv1.txt
and appends AUX-HV1-LEAKSCAN to the ledger. Exit code 1 on a hard fail.
"""
from __future__ import annotations

import io
import json
import re
import sys

import hv1_common as H
from hv1_common import ROOT

POLICY = ["ALWAYS", "MODEL-ALL", "LOOKUP-ALL", "LOOKUP-LIT", "LOOKUP", "MODEL", "ORACLE"]
CONDITION_WORDS = ["retag", "oracle", "gated", "haiku", "gpt", "top-5", "top 5", "model rule", "lookup rule", "literal lookup", "literal rule"]
PLAIN_REVIEW_WORDS = ["literal"]        # the plain word, listed for review when it appears inside a record or answer text
META_WORDS = ["amend", "before", "after", "mini", "detector"]
VERSION_WORDS = ["PRESENT-NEGATIVE", "PRESENT", "ABSENT", "PN", "V2R", "v2", "v3"]
DECIMAL = re.compile(r"(?<![\w.])[01]\.\d{2}(?![\w.])")
VIEW_KEY = re.compile(r"frbank:|cf2?:|\|[0-9a-f]{16}\b")
RECORD_ID = re.compile(r"\blb_\d+|\blb_r82_")
# the content columns per sheet (1-based): question, old value, true now, answer; TRIPLES' question and notes column
CONTENT = {"ANSWERS": {2, 3, 4, 5}, "PRACTICE": {2, 3, 4, 5, 7, 6}, "TRIPLES": {4, 5, 6}, "README": {1}, "ATTESTATION": set()}
# PRACTICE holds two tables; its KEY label column (6 for answers, 5 for stores) legitimately carries the label words; the "Why" column too
PRACTICE_LABEL_COLS = {5, 6, 7}


def personas():
    key = json.load(io.open(H.KEY_DIR / "key_hv1.json", encoding="utf-8"))
    ps = {it["persona"] for it in key["hva"]["items"].values()} | {t["persona"] for t in key["hvb"]["triples"]}
    return sorted(ps), key


def word_hits(text, words, whole=True):
    out = []
    for w in words:
        pat = r"(?<![A-Za-z])" + re.escape(w) + r"(?![A-Za-z])" if whole else re.escape(w)
        if re.search(pat, text, re.IGNORECASE if w.islower() or " " in w else 0):
            out.append(w)
    return out


def scan(path, ann, plist):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    hard, review, n_cells = [], [], 0
    for ws in wb.worksheets:
        content_cols = CONTENT.get(ws.title, set())
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                n_cells += 1; s = str(c.value); where = f"{ws.title}!{c.coordinate}"
                is_content = c.column in content_cols
                hits = {}
                p = word_hits(s, POLICY)                                  # case-sensitive whole words
                if p: hits["policy"] = p
                cw = word_hits(s, CONDITION_WORDS)
                if cw: hits["condition"] = cw
                mw = word_hits(s, META_WORDS)
                if mw: hits["meta"] = mw
                pw = word_hits(s, PLAIN_REVIEW_WORDS)
                if pw: hits["plain"] = pw
                if ws.title == "PRACTICE" and c.column in PRACTICE_LABEL_COLS:
                    pass                                                  # the practice key's label words are meant to be there
                else:
                    vw = [w for w in VERSION_WORDS if re.search(r"(?<![A-Za-z])" + re.escape(w) + r"(?![A-Za-z])", s)]
                    if vw: hits["version"] = vw
                if DECIMAL.search(s): hits["decimal"] = DECIMAL.findall(s)
                if VIEW_KEY.search(s): hits["key"] = VIEW_KEY.findall(s)[:3]
                if RECORD_ID.search(s): hits["record_id"] = RECORD_ID.findall(s)[:3]
                pk = [q for q in plist if re.search(r"(?<![A-Za-z])" + re.escape(q) + r"(?![A-Za-z])", s)]
                if pk and not is_content: hits["persona_key"] = pk       # the names appear inside the content by design (registration 2)
                if not hits:
                    continue
                entry = (ann, where, hits, s[:120].replace("\n", " / "))
                if not is_content:
                    hard.append(entry)
                elif hits.get("policy") or hits.get("condition") or hits.get("key") or hits.get("record_id"):
                    hard.append(entry)
                else:
                    review.append(entry)
    return n_cells, hard, review


def main():
    H.assert_registration()
    plist, key = personas()
    lines, all_hard, all_review, cells = [], [], [], {}
    for ann in ("A1", "A2"):
        path = ROOT / key["workbooks"][ann]["path"].replace(chr(92), "/")     # the key of 19 Sep carries Windows separators
        n, hard, review = scan(path, ann, plist)
        cells[ann] = n; all_hard += hard; all_review += review
        lines.append(f"{ann}: {path.relative_to(ROOT)} sha256 {H.sha_file(path)} cells {n} hard {len(hard)} review {len(review)}")
    lines.append("")
    lines.append("HARD FAILS" if all_hard else "HARD FAILS: none")
    for e in all_hard:
        lines.append(f"  {e[0]} {e[1]} {e[2]} :: {e[3]}")
    lines.append("")
    lines.append(f"CONTENT-COLUMN OCCURRENCES FOR REVIEW ({len(all_review)}): words that appear inside the material being judged")
    summary = {}
    for e in all_review:
        for k, v in e[2].items():
            for w in (v if isinstance(v, list) else [v]):
                summary[(k, str(w))] = summary.get((k, str(w)), 0) + 1
    for (k, w), n in sorted(summary.items(), key=lambda x: -x[1]):
        lines.append(f"  {k}:{w} x{n}")
    for e in all_review[:400]:
        lines.append(f"  {e[0]} {e[1]} {e[2]} :: {e[3]}")
    out = H.RES_DIR / "leakscan_hv1.txt"; out.parent.mkdir(parents=True, exist_ok=True)
    io.open(out, "w", encoding="utf-8", newline="\n").write("\n".join(lines) + "\n")
    verdict = "FAIL" if all_hard else "PASS"
    H.append_ledger("AUX-HV1-LEAKSCAN (Phase 2)",
                    f"`tools/leakscan_hv1.py` over every cell of both blank workbooks ({cells}): **{verdict}**; hard fails {len(all_hard)}; content-column "
                    f"occurrences listed for review {len(all_review)} (by word: " + ", ".join(f"{k}:{w} x{n}" for (k, w), n in sorted(summary.items(), key=lambda x: -x[1])) +
                    f"). Output `human_validation/hv1/results/leakscan_hv1.txt`.")
    print("\n".join(lines[:4])); print(f"leak scan: {verdict}; review list {len(all_review)} -> {out.relative_to(ROOT)}")
    sys.exit(1 if all_hard else 0)


if __name__ == "__main__":
    main()
