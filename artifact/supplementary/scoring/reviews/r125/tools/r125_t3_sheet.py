# -*- coding: utf-8 -*-
"""r125_t3_sheet.py -- Task 3.4: the author's read sheet for the 22 obligations the alignment does not mark. NO MODEL CALL. $0.

One workbook, reviews/r125/build/AUTHOR_SHEET_R125_T3.xlsx, one tab (no 35-row duplicate-audit workbook exists anywhere in the tree, so there
is nothing to add a tab to): the 22 unmarked obligations (AUX-DS-FLAGCENSUS's list, reproduced by r125_t3_frbank.py) and 20 seeded control rows
drawn from the 82 join-marked obligations (seed 42, the sessions' convention; controls whose nearest record within the item's own store variant
is the join-marked one at 0.80 or above), 42 rows shuffled with seed 42. Each row: the persona, the question, the state, the manifest's
forbidding fact and its counterpart (the retraction statement, the new fact, or 'expired'), the nearest stored record within the item's store
variant (the best token coverage of the forbidding fact, the join's own tokeniser and coverage; the detector's discriminative tokens where a
counterpart exists), its overlap score, the store's own lifecycle mark on it; one answer column "is this the forbidding record: YES / NO /
NOT STORED" (data validation) and a notes column, both unlocked; every other cell locked (sheet protection). Nothing on the sheet says which
rows are controls or what the join decided. The key (row -> item id, control flag, nearest entry id) is reviews/r125/build/
AUTHOR_SHEET_R125_T3_key.json; both files' sha256 are recorded in reviews/r125/results/r125_t3_sheet.json and the report.
"""
from __future__ import annotations

import hashlib
import sys

import numpy as np

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
import r125_common as C  # noqa: E402
from r125_common import r4  # noqa: E402
sys.path.insert(0, str(C.ROOT / "resp_review" / "tools")); sys.path.insert(0, str(C.ROOT / "e0_metamemory" / "tools")); sys.path.insert(0, str(C.ROOT / "e1_metamemory" / "tools"))
import resp_common as RC  # noqa: E402
from e0_common import MIN_DISCR, MIN_SHARED, cov, load_joins, toks  # noqa: E402
sys.path.insert(0, str(C.TOOLS))
from r125_t3_frbank import forbidding_events  # noqa: E402

XLSX = C.BUILD / "AUTHOR_SHEET_R125_T3.xlsx"
KEY = C.BUILD / "AUTHOR_SHEET_R125_T3_key.json"
OUT = C.RESULTS / "r125_t3_sheet.json"
SEED = 42
N_CONTROLS = 20
STATE_WORD = {"sup": "supersession", "ret": "retraction", "exp": "expiry"}


def nearest_in(J, fact, counterpart, allowed, members):
    ft = toks(fact, J.drop)
    if counterpart:
        nt = toks(counterpart, J.drop); do = ft - nt
        base = do if len(do) >= MIN_DISCR else ft
    else:
        base = ft
    best = (0.0, None); best_m = (0.0, None)
    for eid in allowed:
        t = J.etoks.get(eid)
        if t is None or len(base & t) < MIN_SHARED:
            continue
        c = cov(base, t)
        if c > best[0]:
            best = (c, eid)
        if eid in members and c > best_m[0]:
            best_m = (c, eid)
    return best, best_m


def own_mark(e):
    tags = []
    if not e.get("is_active"):
        tags.append("superseded (not active)")
    if e.get("is_retracted"):
        tags.append("retracted")
    return "; ".join(tags) if tags else "active"


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    fr = C.jload(C.RESULTS / "r125_t3_frbank.json")
    by_id = {r["item_id"]: r for r in fr["obligations"]}
    sub = {it["item_id"]: it for it in RC.load_cell() if it["y_nok"] == 1}
    joins = load_joins()
    rows = []
    for iid, r in by_id.items():
        it = sub[iid]; J = joins[it["persona"]]; allowed = set(RC.variant_ids(it)); members = set(it["pn_members"])
        best_row = None
        for lab, fact, cp in forbidding_events(J, it["state"]):
            (c, eid), (cm, em) = nearest_in(J, fact, cp, allowed, members)
            cand = {"event": lab, "fact": fact, "counterpart": cp, "entry": eid, "cov": r4(c), "member_cov": r4(cm), "entry_is_member": bool(eid in members)}
            if best_row is None or (cand["member_cov"] or 0) > (best_row["member_cov"] or 0) or ((cand["member_cov"] or 0) == (best_row["member_cov"] or 0) and (cand["cov"] or 0) > (best_row["cov"] or 0)):
                best_row = cand
        e = J.by_id.get(best_row["entry"]) if best_row and best_row["entry"] else None
        rows.append({"item_id": iid, "persona": it["persona"], "question": it["question"], "state": it["state"], "join_marked": r["join_marked_0.80"], "own_marked": r["own_marked"],
                     "event": best_row["event"], "fact": best_row["fact"], "counterpart": best_row["counterpart"] if best_row["counterpart"] else ("(the manifest marks the arrangement expired)" if it["state"] == "exp" else ""),
                     "entry": best_row["entry"], "entry_text": e["text"] if e else "(no stored record shares two content tokens with the fact)", "cov": best_row["cov"], "entry_is_member": best_row["entry_is_member"],
                     "entry_own_mark": own_mark(e) if e else "", "entry_category": e["category"] if e else "", "n_variant_entries": len(allowed)})
    unmarked = [r for r in rows if not r["join_marked"]]; assert len(unmarked) == 22
    pool = [r for r in rows if r["join_marked"] and (r["cov"] or 0) >= 0.80]
    rng = np.random.default_rng(SEED)
    controls = [pool[i] for i in sorted(rng.choice(len(pool), N_CONTROLS, replace=False))]
    sheet = [dict(r, is_control=False) for r in unmarked] + [dict(r, is_control=True) for r in controls]
    order = rng.permutation(len(sheet)); sheet = [sheet[i] for i in order]
    # ---- the workbook
    wb = Workbook(); ws = wb.active; ws.title = "R125_T3_read"
    head = ["row", "persona", "question", "state", "the manifest's forbidding fact", "its counterpart (retraction statement / new fact)", "nearest stored record (within the item's store)",
            "overlap (token coverage)", "the store's own mark on that record", "is this the forbidding record: YES / NO / NOT STORED", "notes"]
    ws.append(head)
    for k, r in enumerate(sheet, 1):
        ws.append([k, r["persona"], r["question"], STATE_WORD[r["state"]], r["fact"], r["counterpart"], r["entry_text"], r["cov"], r["entry_own_mark"], "", ""])
    widths = [5, 11, 44, 12, 44, 44, 52, 9, 22, 28, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.protection = Protection(locked=True)
        row[9].protection = Protection(locked=False); row[10].protection = Protection(locked=False)
        row[9].fill = PatternFill("solid", fgColor="FFF2CC"); row[10].fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"YES,NO,NOT STORED"', allow_blank=True, showErrorMessage=True, errorTitle="Verdict", error="YES, NO or NOT STORED")
    ws.add_data_validation(dv); dv.add(f"J2:J{ws.max_row}")
    ws.freeze_panes = "A2"
    ws.protection.sheet = True; ws.protection.password = "r125"; ws.protection.formatColumns = False; ws.protection.formatRows = False
    C.BUILD.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    key = {"seed": SEED, "n_unmarked": 22, "n_controls": N_CONTROLS, "control_pool": len(pool), "rows": [{"row": k, "item_id": r["item_id"], "is_control": r["is_control"], "join_marked_0.80": r["join_marked"],
                                                                                                           "own_marked": r["own_marked"], "nearest_entry": r["entry"], "cov": r["cov"], "entry_is_member": r["entry_is_member"]} for k, r in enumerate(sheet, 1)]}
    C.jdump(key, KEY)
    res = {"meta": {"script": "reviews/r125/tools/r125_t3_sheet.py", "model_calls": 0}, "xlsx": str(XLSX.relative_to(C.ROOT)), "xlsx_sha256": hashlib.sha256(XLSX.read_bytes()).hexdigest(),
           "key": str(KEY.relative_to(C.ROOT)), "key_sha256": hashlib.sha256(KEY.read_bytes()).hexdigest(), "n_rows": len(sheet), "n_unmarked": 22, "n_controls": N_CONTROLS, "control_pool": len(pool), "seed": SEED,
           "unmarked_nearest": [{"item_id": r["item_id"], "state": r["state"], "nearest_entry": r["entry"], "cov": r["cov"], "entry_is_member": r["entry_is_member"], "own_mark": r["entry_own_mark"]} for r in unmarked],
           "protection": "sheet protection on (password r125); columns J (verdict, data-validated to YES / NO / NOT STORED) and K (notes) unlocked; every other cell locked"}
    C.jdump(res, OUT)
    print(f"sheet {XLSX.relative_to(C.ROOT)} sha256 {res['xlsx_sha256'][:16]}...; key sha256 {res['key_sha256'][:16]}...; {len(sheet)} rows (22 unmarked + {N_CONTROLS} controls from a pool of {len(pool)}), seed {SEED}")
    for r in unmarked:
        print(f"  {r['item_id']:44s} {r['state']} nearest {r['entry']} cov {r['cov']} {'member' if r['entry_is_member'] else 'not a member'} own {r['entry_own_mark']}")


if __name__ == "__main__":
    main()
