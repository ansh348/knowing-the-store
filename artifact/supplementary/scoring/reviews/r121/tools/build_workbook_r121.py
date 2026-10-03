# -*- coding: utf-8 -*-
"""build_workbook_r121.py -- program R121: the author-check workbook for the crossed questions (ruling 5; the reduced design). NO MODEL CALL.

The reduced check design (the author, 2026-09-19): the checked set is all of population B, every population-A row whose
seed-42 question is off the modal template, and a seeded 20 of the conforming A rows (random.Random(42).sample over the
conforming pair_ids sorted); the remaining conforming A rows are the reserve, checked in full only if any sampled row fails.
The modal template of population A is the most frequent form of the seed-42 questions after masking: the first three words,
an ellipsis, the last three words (a question of six words or fewer is its own form).

Tabs: `seed42` (the checked set, the seed-42 question; four YES/NO columns (i)-(iv) and Notes), `seed4242` (the checked set,
the second lineage's question, temperature 0.7, seed 4242; consulted only for rows failed on `seed42`), `reserve_seed42` and
`reserve_seed4242` (the reserve rows, the same layout, consulted only if any sampled row fails), `CRITERIA`. A draft the
automatic gate rejected is shown with its reason and its verdict cells locked (not a candidate). No IDENTICAL greying (ruling
2 on the Phase 1 report): identity between the two lineages is reported in the ledger. The disclosure line (the author's wording of
2026-09-19) heads every tab: "The author's check was assisted by an AI model that pre-flagged candidate failures; every verdict was made by the author and entered from the author's dictated list."

Every non-verdict cell is locked and its text hashed into build/workbook_lock_r121.json. Writes build/AUTHOR_CHECK_R121.xlsx and
prints its sha256 (AUX-CF2-IX-CHECK, "as built").
"""
from __future__ import annotations

import collections
import datetime as dt
import hashlib
import io
import json
import pathlib
import random

ROOT = pathlib.Path(__file__).resolve().parents[3]
CF2 = ROOT / "r82_counterfactual" / "cf2"
BUILD = ROOT / "reviews" / "r121" / "build"
QF = BUILD / "cross_questions.json"
OUT = BUILD / "AUTHOR_CHECK_R121.xlsx"
LOCK = BUILD / "workbook_lock_r121.json"
SAMPLE_SEED, SAMPLE_N = 42, 20

CRITERIA = [
    ("(i) asks about change / history (A) or about the past event (B)", "YES if the question asks how the attribute has changed over time (population A) or asks, in the past tense, what the event was or when it took place (population B)."),
    ("(ii) old content is a proper part of its answer", "YES if, under Definition 1, the superseded value (A) or the expired event's record (B) is properly part of a full answer to this question, so that using it is licensed."),
    ("(iii) names no stored value", "YES if the question states no value of the attribute as the entry text or the PN successor records it. Naming the slot's topic as the attribute path names it is allowed (the CF renderer's rule; ruling 3)."),
    ("(iv) same attribute as the original", "YES if the question concerns the same attribute as the original question."),
]
AI_LINE = "The author's check was assisted by an AI model that pre-flagged candidate failures; every verdict was made by the author and entered from the author's dictated list."

HEAD_42 = [
    ("h", "R121 author check, tab seed42: judge every row on this tab (the checked set: all of population B, every population-A row off the modal template, and a seeded 20 of the conforming A rows)."),
    ("t", "A question passes only on four YES. Mark every one of the four columns for every row (YES or NO); use Notes for anything worth recording. "
          "The tab seed4242 is consulted ONLY for the rows you fail here; never choose between two passing candidates. If any of the 20 sampled conforming A rows fails, the reserve tabs are checked in full."),
    ("t", "Criterion (iii): the question may name the slot's topic as the attribute path names it (an event, a place, a product); it may not state the attribute's VALUE "
          "as the entry text or the PN successor records it. Population A = supersession (CHANGE intent); B = expiry (PAST intent). A row marked GATE-REJECTED is not a candidate; go to the other tab for it."),
    ("t", AI_LINE + " Locked columns are the record; only the four verdict columns and Notes are editable."),
]
HEAD_4242 = [
    ("h", "R121 author check, tab seed4242: consult a row here ONLY if you failed the same row on tab seed42 (or that row was GATE-REJECTED there)."),
    ("t", "This tab carries the registered second lineage: the same instruction and model at temperature 0.7, seed 4242, one question per pair. Judge a row here by "
          "the same four criteria only where the seed42 row failed. A fallback that is word-for-word the seed42 question cannot rescue a failed row (a drop by rule)."),
    ("t", AI_LINE + " Locked columns are the record; only the four verdict columns and Notes are editable."),
]
HEAD_RES_42 = [
    ("h", "R121 author check, tab reserve_seed42: the conforming population-A rows NOT in the seeded sample. Consult this tab ONLY if any sampled row on tab seed42 failed; then judge every row here."),
    ("t", AI_LINE + " Locked columns are the record; only the four verdict columns and Notes are editable."),
]
HEAD_RES_4242 = [
    ("h", "R121 author check, tab reserve_seed4242: the second lineage for the reserve rows. Consult a row here ONLY if the reserve is being checked and you failed the same row on tab reserve_seed42."),
    ("t", AI_LINE + " Locked columns are the record; only the four verdict columns and Notes are editable."),
]


def sha_text(s):
    return hashlib.sha256(str(s).encode("utf-8")).hexdigest()


def template(q):
    w = q.split()
    return " ".join(w[:3]) + " … " + " ".join(w[-3:]) if len(w) > 6 else q


def design(rows):
    """-> (checked pair_ids, reserve pair_ids, the modal template, the off-template A ids, the sampled ids, the form census)."""
    a = [r for r in rows if r["population"] == "A"]
    forms = collections.Counter(template(r["primary"]["question"]) for r in a)
    modal = forms.most_common(1)[0][0]
    off = [r["pair_id"] for r in a if template(r["primary"]["question"]) != modal]
    conforming = sorted(r["pair_id"] for r in a if template(r["primary"]["question"]) == modal)
    sampled = sorted(random.Random(SAMPLE_SEED).sample(conforming, SAMPLE_N))
    checked = {r["pair_id"] for r in rows if r["population"] == "B"} | set(off) | set(sampled)
    reserve = [p for p in conforming if p not in set(sampled)]
    return checked, reserve, modal, off, sampled, forms


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    doc = json.load(io.open(QF, encoding="utf-8")); rows = doc["questions"]; meta = doc["meta"]
    assert len(rows) == 118 and meta.get("instruction_version") == 2 and meta.get("fallback_lineage") == "lineage2"
    pairs = {p["pair_id"]: p for p in json.load(io.open(CF2 / "cf_pairs_cf2.json", encoding="utf-8"))["pairs"]}
    pn = {r["pair_id"]: r for r in json.load(io.open(CF2 / "cf_pn_texts_cf2.json", encoding="utf-8"))["records"]}
    checked, reserve, modal, off, sampled, forms = design(rows)
    index = {r["pair_id"]: n for n, r in enumerate(rows, 1)}

    wb = Workbook(); locked = {}
    wrap = Alignment(wrap_text=True, vertical="top"); bold = Font(bold=True); grey = PatternFill("solid", fgColor="DDDDDD"); unlocked = Protection(locked=False)

    def put(ws, coord, value, lock=True, **fmt):
        c = ws[coord]; c.value = value; c.alignment = wrap
        for k, v in fmt.items():
            setattr(c, k, v)
        if lock:
            if value not in (None, ""):
                locked[f"{ws.title}!{coord}"] = sha_text(value)
        else:
            c.protection = unlocked
        return c

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    cols = ["n", "pair_id", "pop", "persona", "slot path", "original intent", "original question", "crossed intent", "crossed question",
            "stored value (entry text)", "PN successor text (A only)", CRITERIA[0][0], CRITERIA[1][0], CRITERIA[2][0], CRITERIA[3][0], "Notes"]
    cw = [5, 34, 5, 10, 26, 14, 42, 10, 46, 44, 44, 16, 16, 16, 16, 36]
    dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True, showErrorMessage=True)
    wb.remove(wb.active)

    def tab(name, label, head, ids):
        ws = wb.create_sheet(name); widths(ws, cw)
        r = 1
        for kind, text in head:
            put(ws, f"A{r}", text, font=bold if kind == "h" else Font()); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols)); r += 1
        r += 1; hdr = r
        for j, h in enumerate(cols, 1):
            put(ws, f"{get_column_letter(j)}{r}", h, font=bold)
        ws.add_data_validation(dv); r += 1
        for q in rows:
            if q["pair_id"] not in ids:
                continue
            p = pairs[q["pair_id"]]; d = q[label]; rejected = not d["gate_ok"]
            vals = [index[q["pair_id"]], q["pair_id"], q["population"], q["persona"], q["slot"], q["original_intent"], q["original_question"], q["crossed_intent"],
                    d["question"] + ("   [GATE-REJECTED: " + "; ".join(d["gate_reasons"]) + "]" if rejected else ""), p["entry_text"],
                    (pn[q["pair_id"]]["new_text"] if q["population"] == "A" else "")]
            for j, v in enumerate(vals, 1):
                put(ws, f"{get_column_letter(j)}{r}", v)
            for j in range(12, 17):
                if rejected:
                    put(ws, f"{get_column_letter(j)}{r}", "GATE-REJECTED" if j < 16 else None, fill=grey)
                else:
                    put(ws, f"{get_column_letter(j)}{r}", None, lock=False)
                    if j < 16:
                        dv.add(f"{get_column_letter(j)}{r}")
            r += 1
        ws.freeze_panes = f"A{hdr + 1}"; ws.protection.sheet = True
        return r - hdr - 1

    n_checked = tab("seed42", "primary", HEAD_42, checked)
    tab("seed4242", "fallback", HEAD_4242, checked)
    n_reserve = tab("reserve_seed42", "primary", HEAD_RES_42, set(reserve))
    tab("reserve_seed4242", "fallback", HEAD_RES_4242, set(reserve))

    ws = wb.create_sheet("CRITERIA"); widths(ws, [46, 110])
    for i, (name, text) in enumerate(CRITERIA, 1):
        put(ws, f"A{i}", name, font=bold); put(ws, f"B{i}", text)
    put(ws, "A6", "Drop rule", font=bold); put(ws, "B6", "A pair whose seed42 question fails (or was gate-rejected) and whose seed4242 question also fails, was gate-rejected, or is word-for-word the seed42 question, is dropped; no replacement is drawn. Drops are reported with their reasons.")
    put(ws, "A7", "Reduced design", font=bold); put(ws, "B7", f"Checked set: all 30 of population B, the {len(off)} population-A rows off the modal template ('{modal}', {forms[modal]} of 88), and a seeded {SAMPLE_N} of the {forms[modal]} conforming A rows (random.Random({SAMPLE_SEED}).sample over the sorted pair_ids). Reserve: the other {len(reserve)} conforming A rows, checked in full only if any sampled row fails.")
    put(ws, "A8", "Record", font=bold); put(ws, "B8", AI_LINE + " The rows that used the second seed are recorded (AUX-CF2-IX-CHECK).")
    ws.protection.sheet = True

    OUT.parent.mkdir(parents=True, exist_ok=True); wb.save(OUT)
    sha = hashlib.sha256(OUT.read_bytes()).hexdigest()
    rec = {"workbook": str(OUT.relative_to(ROOT)).replace("\\", "/"), "sha256_as_built": sha, "built": dt.datetime.now().isoformat(timespec="seconds"),
           "questions_file": str(QF.relative_to(ROOT)).replace("\\", "/"), "instruction_version": 2, "deviation": "D-IX-1",
           "prompt_sha256": meta["primary"]["prompt_sha256"], "questions_sha256_seed42": meta["primary"]["questions_sha256"],
           "fallback_lineage": {"temperature": meta["lineage2"]["temperature"], "seed": meta["lineage2"]["seed"], "lineage_sha256": meta["lineage2"]["lineage_sha256"], "questions_sha256": meta["lineage2"]["questions_sha256"]},
           "design": {"modal_template": modal, "template_forms": dict(forms), "off_template_A": off, "sample_seed": SAMPLE_SEED, "sample_n": SAMPLE_N, "sampled_A": sampled,
                      "checked": sorted(checked), "reserve": reserve, "n_checked": n_checked, "n_reserve": n_reserve},
           "n_rows": len(rows), "n_identical_fallback": sum(q["fallback_identical"] for q in rows),
           "gate_rejected": {"primary": [q["pair_id"] for q in rows if not q["primary"]["gate_ok"]], "fallback": [q["pair_id"] for q in rows if not q["fallback"]["gate_ok"]]},
           "ai_assisted_line": AI_LINE, "locked_cells": locked}
    io.open(LOCK, "w", encoding="utf-8", newline="\n").write(json.dumps(rec, indent=1, sort_keys=True) + "\n")
    print(f"{OUT.relative_to(ROOT)}  sha256 {sha}\n  checked {n_checked} (B 30 + off-template A {len(off)} + sampled A {SAMPLE_N}); reserve {n_reserve}; modal template '{modal}' ({forms[modal]} of 88)\n"
          f"  locked cells {len(locked)}; identical fallbacks {rec['n_identical_fallback']}; gate-rejected primary {rec['gate_rejected']['primary']}, fallback {rec['gate_rejected']['fallback']}")


if __name__ == "__main__":
    main()
