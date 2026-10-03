# -*- coding: utf-8 -*-
"""enter_verdicts_r121.py -- program R121: enter the author's dictated verdicts (2026-09-19) into the as-built workbook, verify the
locked cells against the lock file, record the returned workbook's sha256, and derive the passing questions, the drops and the
second-lineage rows. NO MODEL CALL.

The verdicts below are the author's, transcribed from the dictated list verbatim (row numbers are the workbook's `n` column, the
index into build/cross_questions.json). Every checked row not listed is four YES on tab seed42. The reserve tabs stay unjudged
(no sampled row failed) and their rows pass under the reduced design.
Writes build/AUTHOR_CHECK_R121_returned.xlsx, build/verdicts_r121.json.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import pathlib

ROOT = pathlib.Path(__file__).resolve().parents[3]
BUILD = ROOT / "reviews" / "r121" / "build"
SENT = BUILD / "AUTHOR_CHECK_R121.xlsx"
LOCK = BUILD / "workbook_lock_r121.json"
RET = BUILD / "AUTHOR_CHECK_R121_returned.xlsx"
OUT = BUILD / "verdicts_r121.json"
QF = BUILD / "cross_questions.json"

YES4 = {"i": "YES", "ii": "YES", "iii": "YES", "iv": "YES"}
# the author's dictated list, 2026-09-19 (row n -> tab -> verdicts and note)
DICTATED = {
    28: {"seed42": ({"i": "YES", "ii": "NO", "iii": "YES", "iv": "YES"}, "record says only that Greg was accepted into a show; nothing about a relationship")},
    33: {"seed42": ({"i": "YES", "ii": "NO", "iii": "YES", "iv": "NO"}, "record is a delivery date, not use; question drifts to usage"),
         "seed4242": ({"i": "YES", "ii": "NO", "iii": "YES", "iv": "NO"}, "record is a delivery date, not use; question drifts to usage")},
    52: {"seed42": ({"i": "YES", "ii": "NO", "iii": "YES", "iv": "YES"}, "asks how the trip went; the store records no outcome")},
    77: {"seed42": ({"i": "YES", "ii": "NO", "iii": "YES", "iv": "NO"}, "record is a concern about one ceremony's logistics; question asks what was used, for events in general")},
    104: {"seed42": ({"i": "YES", "ii": "YES", "iii": "YES", "iv": "NO"}, "slot name read as a time reference; the attribute is talking about old times"),
          "seed4242": ({"i": "YES", "ii": "YES", "iii": "YES", "iv": "NO"}, "slot name read as a time reference; the attribute is talking about old times")},
    27: {"seed42": ({"i": "YES", "ii": "NO", "iii": "YES", "iv": "YES"}, "record has no 'when'"),
         "seed4242": (dict(YES4), None)},
    32: {"seed42": ("GATE-REJECTED", "the seed42 row is GATE-REJECTED; marked failed"),
         "seed4242": (dict(YES4), None)},
    21: {"seed42": (dict(YES4), "'in the past' is ambiguous, but the planned count is a proper part of the answer")},
}
for n in (6, 19, 85, 99, 109, 116):
    DICTATED[n] = {"seed42": (dict(YES4), "slot key carries the old value; passes (iii) as registered")}
EXPECTED = {"A_pass": 87, "A_total": 88, "B_pass": 26, "B_total": 30, "drops": 5, "second_lineage": 2}
COLS = {"i": 12, "ii": 13, "iii": 14, "iv": 15, "note": 16}


def sha_text(s):
    return hashlib.sha256(str(s).encode("utf-8")).hexdigest()


def main():
    from openpyxl import load_workbook
    lock = json.load(io.open(LOCK, encoding="utf-8"))
    sent_sha = hashlib.sha256(SENT.read_bytes()).hexdigest()
    assert sent_sha == lock["sha256_as_built"], ("the sent workbook is not the one the lock describes", sent_sha, lock["sha256_as_built"])
    rows = json.load(io.open(QF, encoding="utf-8"))["questions"]
    wb = load_workbook(SENT)
    entered = {}

    def hdr_row(ws):
        return [r for r in range(1, 10) if ws.cell(r, 1).value == "n"][0]

    # tab seed42: every checked row four YES unless dictated; tab seed4242: only the dictated rows
    for tab in ("seed42", "seed4242"):
        ws = wb[tab]; h = hdr_row(ws)
        for r in range(h + 1, ws.max_row + 1):
            n = ws.cell(r, 1).value
            if n is None:
                continue
            n = int(n)
            d = DICTATED.get(n, {})
            if tab == "seed42":
                verdict, note = d.get("seed42", (dict(YES4), None))
            else:
                if "seed4242" not in d:
                    continue
                verdict, note = d["seed4242"]
            if verdict == "GATE-REJECTED":
                assert ws.cell(r, COLS["i"]).value == "GATE-REJECTED", (tab, n)
                ws.cell(r, COLS["note"]).value = note
                entered[(tab, n)] = {"verdict": "GATE-REJECTED", "note": note}
                continue
            for k, c in COLS.items():
                if k == "note":
                    ws.cell(r, c).value = note
                else:
                    assert ws.cell(r, c).value is None and not ws.cell(r, c).protection.locked, (tab, n, k)
                    ws.cell(r, c).value = verdict[k]
            entered[(tab, n)] = {"verdict": verdict, "note": note}
    wb.save(RET)

    # verify every locked cell of the returned file against the lock
    wb2 = load_workbook(RET); bad = []
    for key, h in lock["locked_cells"].items():
        tab, coord = key.split("!")
        v = wb2[tab][coord].value
        if sha_text(v) != h:
            bad.append(key)
    assert not bad, ("locked cells changed", bad[:10])
    ret_sha = hashlib.sha256(RET.read_bytes()).hexdigest()

    # derive the outcome per pair
    checked = set(lock["design"]["checked"]); reserve = set(lock["design"]["reserve"])
    out = {}; drops = []; second = []
    for n, q in enumerate(rows, 1):
        pid = q["pair_id"]
        if pid in reserve:
            out[pid] = {"n": n, "population": q["population"], "status": "PASS-RESERVE-UNJUDGED", "question": q["primary"]["question"], "lineage": "primary"}
            continue
        assert pid in checked
        v42 = entered[("seed42", n)]
        p42 = v42["verdict"] != "GATE-REJECTED" and all(v42["verdict"][k] == "YES" for k in ("i", "ii", "iii", "iv"))
        if p42:
            out[pid] = {"n": n, "population": q["population"], "status": "PASS", "question": q["primary"]["question"], "lineage": "primary", "note": v42["note"]}
            continue
        fail42 = "GATE-REJECTED" if v42["verdict"] == "GATE-REJECTED" else "NO on " + ", ".join(k for k in ("i", "ii", "iii", "iv") if v42["verdict"][k] == "NO")
        if q["fallback_identical"]:
            drops.append(pid); out[pid] = {"n": n, "population": q["population"], "status": "DROP", "reason": f"seed42 {fail42}; fallback byte-identical to the primary", "note": v42["note"]}
            continue
        if not q["fallback"]["gate_ok"]:
            drops.append(pid); out[pid] = {"n": n, "population": q["population"], "status": "DROP", "reason": f"seed42 {fail42}; fallback gate-rejected", "note": v42["note"]}
            continue
        v44 = entered.get(("seed4242", n))
        assert v44 is not None, ("a failed seed42 row without a seed4242 verdict", n)
        p44 = all(v44["verdict"][k] == "YES" for k in ("i", "ii", "iii", "iv"))
        if p44:
            second.append(pid); out[pid] = {"n": n, "population": q["population"], "status": "PASS-SECOND-LINEAGE", "question": q["fallback"]["question"], "lineage": "lineage2",
                                            "seed42_failure": fail42, "note": v42["note"]}
        else:
            fail44 = "NO on " + ", ".join(k for k in ("i", "ii", "iii", "iv") if v44["verdict"][k] == "NO")
            drops.append(pid); out[pid] = {"n": n, "population": q["population"], "status": "DROP", "reason": f"seed42 {fail42}; seed4242 {fail44}", "note": v44["note"] or v42["note"]}
    counts = {"A_pass": sum(1 for v in out.values() if v["population"] == "A" and v["status"] != "DROP"), "A_total": 88,
              "B_pass": sum(1 for v in out.values() if v["population"] == "B" and v["status"] != "DROP"), "B_total": 30,
              "drops": len(drops), "second_lineage": len(second),
              "checked": len(checked), "reserve_unjudged": len(reserve), "pass_primary_checked": sum(1 for v in out.values() if v["status"] == "PASS")}
    ok = all(counts[k] == EXPECTED[k] for k in EXPECTED)
    rec = {"sent_sha256": sent_sha, "returned_sha256": ret_sha, "returned": str(RET.relative_to(ROOT)).replace("\\", "/"), "entered": dt.datetime.now().isoformat(timespec="seconds"),
           "locked_cells_verified": len(lock["locked_cells"]), "disclosure": lock["ai_assisted_line"], "counts": counts, "expected": EXPECTED, "counts_match_expected": ok,
           "drops": [{"pair_id": p, **{k: out[p][k] for k in ("n", "population", "reason", "note")}} for p in drops],
           "second_lineage": [{"pair_id": p, **{k: out[p][k] for k in ("n", "population", "question", "seed42_failure", "note")}} for p in second],
           "notes_on_passing_rows": {n: DICTATED[n]["seed42"][1] for n in sorted(DICTATED) if DICTATED[n]["seed42"][0] == YES4 and DICTATED[n]["seed42"][1]},
           "pairs": out}
    io.open(OUT, "w", encoding="utf-8", newline="\n").write(json.dumps(rec, indent=1, sort_keys=True, ensure_ascii=False) + "\n")
    print(f"entered {len(entered)} verdict rows; locked cells verified {len(lock['locked_cells'])}; returned sha256 {ret_sha}")
    print("counts", json.dumps(counts), "expected", json.dumps(EXPECTED), "MATCH" if ok else "MISMATCH -- STOP")
    for d in rec["drops"]:
        print("  DROP", d["n"], d["pair_id"], "|", d["reason"])
    for s in rec["second_lineage"]:
        print("  SECOND", s["n"], s["pair_id"], "|", s["question"])
    if not ok:
        raise SystemExit("counts do not match the author's expected result")


if __name__ == "__main__":
    main()
